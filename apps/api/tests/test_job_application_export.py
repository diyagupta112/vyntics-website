"""Real workbook content and export HTTP boundary tests."""

import asyncio
from io import BytesIO
from types import SimpleNamespace
from unittest.mock import AsyncMock, patch

import pytest
from openpyxl import load_workbook
from sqlalchemy.exc import SQLAlchemyError
from sqlalchemy.ext.asyncio import AsyncSession

from app.api.dependencies.job_application_export import get_job_application_export_service
from app.services.job_application_export import (
    ApplicationExport, EXCEL_MEDIA_TYPE, HEADERS, ExcelGenerationError,
    JobApplicationExportService, _workbook,
)
from app.services.job_applications import JobApplicationCareerNotFoundError
from tests.auth_helpers import TEST_ADMIN
from tests.test_auth import _admin, authenticated_api
from tests.test_job_application_service import CAREER_ID, _application


def _service():
    session = AsyncMock(spec=AsyncSession)
    service = JobApplicationExportService(session)
    service._applications = AsyncMock()
    service._careers = AsyncMock()
    service._audits = AsyncMock()
    return service, session


def test_all_export_includes_historical_rows_values_and_safe_audit():
    service, session = _service()
    service._applications.list_all.return_value = [
        _application(name="=HYPERLINK(\"evil\")", currently_working=True),
        _application(career_id=None, career_title_snapshot="Deleted Career",
                     currently_working=False, resume_url=None, notes="Private note"),
    ]
    exported = asyncio.run(service.export(None, actor=TEST_ADMIN))
    workbook = load_workbook(BytesIO(exported.content))
    sheet = workbook.active
    assert exported.filename == "Vyntics_Job_Applications_All.xlsx"
    assert tuple(cell.value for cell in sheet[1]) == HEADERS
    assert sheet.max_row == 3
    assert sheet["A2"].data_type == "s"
    assert sheet["D3"].value == "Deleted Career"
    assert sheet["E2"].value == "4 years 6 months"
    assert sheet["F2"].value == "Yes" and sheet["F3"].value == "No"
    assert sheet["H2"].value == "30 Days"
    assert sheet["I2"].value == "New"
    assert sheet["J2"].value == "Available in Admin Panel"
    assert sheet["J3"].value == "Not provided"
    assert sheet["K3"].value == "Private note"
    assert sheet["L2"].value.endswith(" UTC")
    assert sheet.freeze_panes == "A2" and sheet.auto_filter.ref == "A1:L3"
    service._careers.get_by_id.assert_not_awaited()
    audit = service._audits.add.await_args.args[0]
    assert audit.action == "export" and audit.resource_type == "job_application"
    assert audit.actor_id == TEST_ADMIN.admin_id
    assert audit.context == {"scope": "all", "career_id": None, "application_count": 2}
    assert "ada" not in str(audit.context).lower()
    session.commit.assert_awaited_once()


def test_career_export_uses_only_filtered_query_and_safe_filename():
    service, _ = _service()
    service._careers.get_by_id.return_value = SimpleNamespace(title='Engineer / "Remote"')
    service._applications.list_for_career.return_value = [_application()]
    exported = asyncio.run(service.export(CAREER_ID, actor=TEST_ADMIN))
    assert exported.filename == "Vyntics_Job_Applications_Engineer_Remote.xlsx"
    service._applications.list_for_career.assert_awaited_once_with(CAREER_ID)
    service._applications.list_all.assert_not_awaited()
    assert service._audits.add.await_args.args[0].context["career_id"] == str(CAREER_ID)


def test_empty_career_export_produces_headers_only():
    service, _ = _service()
    service._careers.get_by_id.return_value = SimpleNamespace(title="Empty")
    service._applications.list_for_career.return_value = []
    exported = asyncio.run(service.export(CAREER_ID, actor=TEST_ADMIN))
    sheet = load_workbook(BytesIO(exported.content)).active
    assert sheet.max_row == 1
    assert tuple(cell.value for cell in sheet[1]) == HEADERS


def test_missing_career_and_generation_failure_do_not_audit():
    service, session = _service()
    service._careers.get_by_id.return_value = None
    with pytest.raises(JobApplicationCareerNotFoundError):
        asyncio.run(service.export(CAREER_ID, actor=TEST_ADMIN))
    service._applications.list_all.return_value = []
    with patch("app.services.job_application_export._workbook", side_effect=ValueError("internals")):
        with pytest.raises(ExcelGenerationError):
            asyncio.run(service.export(None, actor=TEST_ADMIN))
    service._audits.add.assert_not_awaited()
    session.commit.assert_not_awaited()


def test_audit_commit_failure_rolls_back():
    service, session = _service()
    service._applications.list_all.return_value = []
    session.commit.side_effect = SQLAlchemyError("internals")
    with pytest.raises(SQLAlchemyError):
        asyncio.run(service.export(None, actor=TEST_ADMIN))
    session.rollback.assert_awaited_once()


@pytest.mark.parametrize("role", ["admin", "superadmin"])
def test_export_http_accepts_both_roles_and_returns_real_xlsx(authenticated_api, role):
    client, _, session, _ = authenticated_api
    session.scalar.return_value = _admin(role=role)
    service = AsyncMock(spec=JobApplicationExportService)
    service.export.return_value = ApplicationExport(_workbook([]), "Vyntics_Job_Applications_All.xlsx")
    client.app.dependency_overrides[get_job_application_export_service] = lambda: service
    response = client.get("/admin/job-applications/export", headers={"Authorization": "Bearer valid"})
    assert response.status_code == 200
    assert response.headers["content-type"] == EXCEL_MEDIA_TYPE
    assert response.headers["content-disposition"] == 'attachment; filename="Vyntics_Job_Applications_All.xlsx"'
    assert response.headers["cache-control"] == "no-store"
    assert load_workbook(BytesIO(response.content)).active.max_row == 1


def test_export_auth_and_openapi(authenticated_api):
    client, _, _, _ = authenticated_api
    assert client.get("/admin/job-applications/export").status_code == 401
    document = client.get("/openapi.json").json()
    operation = document["paths"]["/admin/job-applications/export"]["get"]
    assert operation["security"] == [{"HTTPBearer": []}]
    assert operation["parameters"][0]["name"] == "career_id"
    assert EXCEL_MEDIA_TYPE in operation["responses"]["200"]["content"]


@pytest.mark.parametrize("error,status", [
    (JobApplicationCareerNotFoundError(), 404),
    (SQLAlchemyError("secret"), 503),
    (ExcelGenerationError("secret"), 500),
])
def test_export_http_sanitizes_failures(authenticated_api, error, status):
    client, _, _, _ = authenticated_api
    service = AsyncMock(spec=JobApplicationExportService)
    service.export.side_effect = error
    client.app.dependency_overrides[get_job_application_export_service] = lambda: service
    response = client.get("/admin/job-applications/export", headers={"Authorization": "Bearer valid"})
    assert response.status_code == status
    assert "secret" not in response.text
